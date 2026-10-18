#!/usr/bin/env python3
"""Download real client-generated reports with synthetic transport, then parse them.
No production requests, real accounts, or attendee data are used. Requires existing
Python Playwright, Chromium, openpyxl and pypdf. Files remain in a temporary directory.
"""
import argparse
import base64
import csv
import io
import json
import os
import re
import tempfile
import subprocess
import warnings
from pathlib import Path
from playwright.sync_api import sync_playwright
from openpyxl import load_workbook
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / 'app/assets'
STAMP = '2026-10-18T12:00:00Z'
PROFILE = dict(role='checker', checkerId='22222222-2222-4222-8222-222222222222', name='Synthetic Faculty', method='google', active=True, reports=True, scanner=False, helpdesk=False, nameConfirmed=True)
SUMMARY = dict(registered=3, eligible=2, blocked=1, attended=2, registeredAttended=1, organizerConfirmed=1, noShow=1, serverIssues=2, issueTickets=1, duplicateAttempts=1, pendingHelpDesk=0, generatedAt=STAMP)
NAME = '=1+1'
ROW = {'Ticket ID': 'DEMO0001', 'Name': NAME, 'Name status': 'Recorded', 'In registration master': True, 'Eligible': True, 'Amount INR': 100, 'Attendance': 'Admitted', 'Admission basis': 'server_validated', 'Note': 'Synthetic किरण <&> "quoted"\nsecond line\ufffe', 'Verified entry UTC': STAMP, 'Issue codes': 'blocked', 'Issue attempts': 1, 'First issue UTC': STAMP, 'Outcome': 'Resolved with validated admission', 'Duplicate attempts': 1, 'First duplicate UTC': STAMP, 'Last duplicate UTC': STAMP, 'Hour IST': '2026-10-18 17:00', 'Record source': NAME, 'Time basis': 'Server-validated entry', 'Admissions': 1, 'Status': 'pending', 'Created UTC': STAMP, 'Recorded by': 'Synthetic checker', 'private_email': 'must-not-export@example.invalid'}


def verify(data, kind, fmt, empty):
    if fmt == 'csv':
        rows = list(csv.reader(io.StringIO(data.decode('utf-8-sig'))))
        assert len(rows) == (1 if empty else 2), (kind, fmt, rows)
        assert 'Report generated UTC' in rows[0]
        assert 'private_email' not in rows[0]
        assert ('Hour IST' if kind == 'activity' else 'Ticket ID') in rows[0]
        if not empty:
            parsed = dict(zip(rows[0], rows[1]))
            assert parsed.get('Record source' if kind == 'activity' else 'Name') == "'" + NAME
    else:
        with warnings.catch_warnings():
            warnings.simplefilter('error')
            workbook = load_workbook(io.BytesIO(data), data_only=False)
        sheet = workbook.active
        rows = list(sheet.values)
        assert len(rows) == (1 if empty else 2), (kind, fmt, rows)
        assert 'Report generated UTC' in rows[0]
        assert 'private_email' not in rows[0]
        if not empty:
            parsed = dict(zip(rows[0], rows[1]))
            column = 'Record source' if kind == 'activity' else 'Name'
            assert parsed[column] == NAME
            assert sheet.cell(2, rows[0].index(column) + 1).data_type == 's'
            if 'Note' in parsed:
                assert 'किरण <&> "quoted"\nsecond line' in parsed['Note']
                assert '\ufffe' not in parsed['Note']
        workbook.close()
    assert b'must-not-export' not in data


def format_only(output):
    script = r"""
const fs=require('node:fs'),{JSDOM}=require('jsdom');const f=JSON.parse(process.argv[1]);
const dom=new JSDOM(fs.readFileSync('app/assets/index.html','utf8'),{runScripts:'outside-only',url:'https://fixture.invalid/'}),w=dom.window,saved=[];let empty=false;
w.TextEncoder=TextEncoder;w.GateClient={native:false,profile:()=>f.profile,on(){},api:async(action,payload)=>{if(action==='me')return f.profile;if(action==='reports-summary')return{stats:f.summary};if(action==='reports-export')return{kind:payload.kind,rows:empty?[]:[f.row],count:empty?0:1,generatedAt:f.summary.generatedAt};throw Error('Unexpected fixture action '+action);},saveReport:(name,mime,bytes)=>saved.push({name,empty,content:Buffer.from(bytes).toString('base64')})};
for(const file of ['report-export.js','reports.js'])w.eval(fs.readFileSync('app/assets/'+file,'utf8'));
(async()=>{await w.NavrangReports.refresh();const select=w.document.getElementById('report-kind');for(empty of [false,true])for(const option of select.options){select.value=option.value;for(const format of ['csv','xlsx'])await w.document.getElementById('report-'+format).onclick();}await w.document.getElementById('report-pdf').onclick();process.stdout.write(JSON.stringify(saved));dom.window.close()})().catch(e=>{console.error(e);process.exitCode=1});
"""
    files=json.loads(subprocess.check_output(['node','-e',script,json.dumps(dict(profile=PROFILE,summary=SUMMARY,row=ROW))],cwd=ROOT))
    assert len(files)==49
    for item in files:
        data=base64.b64decode(item['content']);name=item['name'];fmt=name.rsplit('.',1)[1]
        if fmt=='pdf':
            text=PdfReader(io.BytesIO(data)).pages[0].extract_text();assert 'Total admitted: 2' in text and 'DEMO0001' not in text
        else:
            kind=name.split('-')[1];verify(data,kind,fmt,item['empty'])
        if output:
            dest=output/('empty' if item['empty'] else 'populated')/name;dest.parent.mkdir(parents=True,exist_ok=True);dest.write_bytes(data)
    return dict(checks=49,datasets=12,actual_browser_downloads=0,parsed_files=49,production_access=False)


def main():
    parser=argparse.ArgumentParser();parser.add_argument('--output',type=Path);parser.add_argument('--formats-only',action='store_true');args=parser.parse_args()
    if args.formats_only:
        result=format_only(args.output)
        if args.output:
            args.output.mkdir(parents=True,exist_ok=True);(args.output/'format-results.json').write_text(json.dumps(result,indent=2))
        print(json.dumps(result));return
    markup = (ASSETS / 'index.html').read_text()
    # Like the existing responsive source QA, this wrapper is in-memory only.
    markup = re.sub(r'<meta[^>]*http-equiv="Content-Security-Policy"[^>]*>', '', markup)
    markup = re.sub(r'<script.*?</script>|<link[^>]*>', '', markup, flags=re.S)
    errors = []
    console_messages = []
    completed = []
    count = 0
    with tempfile.TemporaryDirectory(prefix='navrang-report-downloads-') as folder, sync_playwright() as pw:
        browser = pw.chromium.launch(headless=True, executable_path=os.getenv('NAVRANG_CHROMIUM', '/usr/bin/chromium'), args=['--no-sandbox', '--disable-dev-shm-usage'])
        context = browser.new_context(accept_downloads=True)
        context.route('**/*', lambda route: route.abort())
        page = context.new_page()
        page.on('pageerror', lambda error: errors.append(str(error)))
        page.on('console', lambda message: console_messages.append(dict(type=message.type,text=message.text)))
        page.set_content(markup)
        page.evaluate('''f=>{
            const storage=new Map();Object.defineProperty(window,'sessionStorage',{value:{getItem:k=>storage.get(k)||null,setItem:(k,v)=>storage.set(k,v),removeItem:k=>storage.delete(k)}});
            sessionStorage.setItem('navrang-session',JSON.stringify({token:'a'.repeat(64),profile:f.profile}));
            window.fixtureEmpty=false;window.fixtureCalls=[];
            window.fetch=async(url,options)=>{
                const {action,payload}=JSON.parse(options.body);fixtureCalls.push({action,payload});let data;
                if(action==='me')data=f.profile;
                else if(action==='reports-summary')data={stats:f.summary};
                else if(action==='reports-export')data={kind:payload.kind,rows:fixtureEmpty?[]:[f.row],count:fixtureEmpty?0:1,generatedAt:f.summary.generatedAt};
                else throw Error('Unexpected fixture action: '+action);
                return new Response(JSON.stringify({ok:true,...data}),{status:200});
            };
        }''', dict(profile=PROFILE, summary=SUMMARY, row=ROW))
        for script in ['live-core.js', 'config.js', 'client.js', 'report-export.js', 'reports.js']:
            page.evaluate((ASSETS / script).read_text())
        # This focused harness does not load gate.js; expose its synthetic faculty
        # destination so ordinary, actionable pointer clicks can be used.
        page.evaluate("document.querySelectorAll('.view').forEach(e=>e.hidden=e.id!=='view-reports');document.getElementById('intro').hidden=true")
        page.evaluate('NavrangReports.refresh()')
        def receive(kind,fmt,empty):
            context_info=dict(kind=kind,format=fmt,empty=empty)
            try:
                # Chromium LocalFrame::ShouldThrottleDownload permits ten burst
                # requests per second. Use a realistic pace, never disable that guard.
                # https://github.com/chromium/chromium/blob/main/third_party/blink/renderer/core/frame/local_frame.cc
                page.wait_for_timeout(500)
                with page.expect_download() as waiting:
                    page.locator('#report-'+fmt).click()
                download=waiting.value
                assert download.failure() is None
                prefix='Navrang-Attendance-Summary-' if fmt=='pdf' else 'Navrang-'+kind+'-'
                assert download.suggested_filename.startswith(prefix)
                destination=(args.output or Path(folder))/('empty' if empty else 'populated')/download.suggested_filename
                destination.parent.mkdir(parents=True,exist_ok=True)
                download.save_as(destination)
                if fmt!='pdf':verify(destination.read_bytes(),kind,fmt,empty)
                else:
                    pdf=PdfReader(destination);assert len(pdf.pages)==1
                    text=pdf.pages[0].extract_text()
                    for expected in ['Total admitted: 2','From registration master: 1','Organizer-confirmed admissions: 1','Exact admission time is unknown',STAMP]:assert expected in text,expected
                    assert 'DEMO0001' not in text
                completed.append(dict(**context_info,filename=download.suggested_filename,bytes=destination.stat().st_size))
                print(json.dumps(dict(download=completed[-1],completed=len(completed))),flush=True)
            except Exception as error:
                target=args.output or Path(folder);target.mkdir(parents=True,exist_ok=True)
                details=dict(**context_info,error=str(error),completed=completed,browser=browser.version,page_errors=errors,console=console_messages,
                    status=page.locator('#report-status').text_content(),selected=page.locator('#report-kind').input_value(),
                    button=dict(visible=page.locator('#report-'+fmt).is_visible(),disabled=page.locator('#report-'+fmt).is_disabled()),
                    calls=page.evaluate('fixtureCalls'))
                (target/'failure.json').write_text(json.dumps(details,indent=2))
                page.screenshot(path=str(target/'failure.png'),full_page=True)
                raise
        kinds = page.eval_on_selector_all('#report-kind option', 'options=>options.map(o=>o.value)')
        assert len(kinds) == 12
        for empty in [False, True]:
            page.evaluate('value=>fixtureEmpty=value', empty)
            for kind in kinds:
                page.locator('#report-kind').select_option(kind)
                for fmt in ['csv', 'xlsx']:
                    receive(kind,fmt,empty)
                    count+=1
        receive('summary','pdf',False)
        count+=1
        assert not errors, errors
        assert page.evaluate("fixtureCalls.filter(x=>x.action==='reports-export').length") == 48
        # Native bridge gets byte-identical Base64 for every format, rather than text corruption.
        page.evaluate('''profile=>{window.Native={local:()=>JSON.stringify({ok:true,data:{profile}}),saveReport:(name,mime,content)=>{window.nativeFile={name,mime,content}}}}''', PROFILE)
        page.evaluate((ASSETS / 'client.js').read_text())
        for fmt in ['csv', 'xlsx', 'pdf']:
            page.evaluate('''fmt=>{const E=NavrangReportExport;const bytes=fmt==='csv'?new TextEncoder().encode(E.csv([{Name:'किरण'}])):fmt==='xlsx'?E.xlsx([{Name:'किरण'}]):E.pdf({registered:0,eligible:0,blocked:0,attended:0,noShow:0,serverIssues:0,issueTickets:0,duplicateAttempts:0,pendingHelpDesk:0});GateClient.saveReport('Navrang-synthetic.'+fmt,fmt==='pdf'?'application/pdf':fmt==='csv'?'text/csv':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',bytes);window.expectedBytes=Array.from(bytes)}''', fmt)
            value = page.evaluate('nativeFile')
            decoded = base64.b64decode(value['content'], validate=True)
            assert decoded == bytes(page.evaluate('expectedBytes'))
            count += 1
        browser.close()
    result=dict(checks=count, datasets=12, browser_downloads=49, native_bridge_formats=3, page_errors=errors, production_access=False)
    if args.output: (args.output/'browser-results.json').write_text(json.dumps(result,indent=2))
    print(json.dumps(result))


if __name__ == '__main__':
    main()
